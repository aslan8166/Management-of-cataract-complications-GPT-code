#!/usr/bin/env python3
"""Generate a fully editable PowerPoint for the Light Adjustable Lens talk.

Unlike the Slidev image-export, every text box, bullet and table cell here is
native PowerPoint content, so it can be edited directly.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---- palette -------------------------------------------------------------
NAVY  = RGBColor(0x1E, 0x3A, 0x8A)
TEAL  = RGBColor(0x0D, 0x94, 0x88)
DTEAL = RGBColor(0x0F, 0x76, 0x6E)
DARK  = RGBColor(0x33, 0x33, 0x33)
GRAY  = RGBColor(0x66, 0x66, 0x66)
AMBER = RGBColor(0xB4, 0x53, 0x09)
LIGHT = RGBColor(0xF1, 0xF5, 0xF9)
TEALBG = RGBColor(0xE6, 0xF4, 0xF2)
AMBBG  = RGBColor(0xFD, 0xF3, 0xE7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
FONT = "Calibri"


def slide():
    return prs.slides.add_slide(BLANK)


def rect(s, x, y, w, h, fill=None, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1)
    return sp


def textbox(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return tb, tf


def setrun(r, text, size, color=DARK, bold=False, italic=False):
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = FONT


def para(tf, text, size=18, color=DARK, bold=False, italic=False, bullet=False,
         level=0, space_after=6, align=PP_ALIGN.LEFT, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.level = level
    p.space_after = Pt(space_after)
    setrun(p.add_run(), text, size, color, bold, italic)
    if not bullet:
        # suppress default bullet
        pPr = p._pPr if p._pPr is not None else p.get_or_add_pPr()
        from pptx.oxml.ns import qn
        for tag in ('a:buChar', 'a:buAutoNum'):
            for e in pPr.findall(qn(tag)):
                pPr.remove(e)
        buNone = pPr.makeelement(qn('a:buNone'), {})
        pPr.append(buNone)
    return p


def header(s, title, kicker=None):
    rect(s, 0, 0, SW, Inches(1.15), fill=NAVY)
    rect(s, 0, Inches(1.15), SW, Pt(4), fill=TEAL)
    tb, tf = textbox(s, Inches(0.6), Inches(0.12), Inches(12.1), Inches(0.95),
                     anchor=MSO_ANCHOR.MIDDLE)
    para(tf, title, 30, WHITE, bold=True, first=True)
    if kicker:
        para(tf, kicker, 13, RGBColor(0xCB, 0xE7, 0xE3), italic=True, space_after=0)


def bullets(tf, items, size=18, gap=8, first=True):
    for i, it in enumerate(items):
        lvl = 0
        txt = it
        bold = False
        if isinstance(it, tuple):
            txt, lvl, bold = it[0], it[1], (it[2] if len(it) > 2 else False)
        prefix = "•  " if lvl == 0 else "      –  "
        para(tf, prefix + txt, size,
             DARK if lvl == 0 else GRAY, bold=bold,
             first=(first and i == 0), space_after=gap, level=0)


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


def card(s, x, y, w, h, fill, title, lines, tcolor=DTEAL):
    rect(s, x, y, w, h, fill=fill)
    tb, tf = textbox(s, x + Inches(0.18), y + Inches(0.12),
                     w - Inches(0.36), h - Inches(0.24))
    para(tf, title, 15, tcolor, bold=True, first=True, space_after=4)
    for ln in lines:
        para(tf, ln, 13, DARK, space_after=3)


# ========================================================================
# 1 — Title
# ========================================================================
s = slide()
rect(s, 0, 0, SW, SH, fill=NAVY)
rect(s, 0, Inches(4.05), SW, Pt(4), fill=TEAL)
tb, tf = textbox(s, Inches(1.0), Inches(2.1), Inches(11.3), Inches(2.0))
para(tf, "The Light Adjustable Lens (LAL)", 46, WHITE, bold=True, first=True,
     space_after=6)
para(tf, "Customizing Vision After Cataract Surgery", 26,
     RGBColor(0x9D, 0xD9, 0xD2), space_after=0)
tb, tf = textbox(s, Inches(1.0), Inches(4.3), Inches(11.3), Inches(1.0))
para(tf, "A non-invasive approach to correcting refractive error after the eye has healed",
     18, RGBColor(0xCB, 0xE7, 0xE3), italic=True, first=True)
notes(s, "Opening (~8-minute talk). The LAL is the first IOL whose power we can "
         "fine-tune AFTER surgery, non-invasively. I'll cover the problem it "
         "solves, how it works, who it's for, and the clinical data.")

# ========================================================================
# 2 — The Clinical Problem
# ========================================================================
s = slide()
header(s, "The Clinical Problem")
tb, tf = textbox(s, Inches(0.6), Inches(1.45), Inches(12.1), Inches(0.7))
para(tf, "Residual refractive error is one of the most important factors "
         "determining visual acuity after cataract surgery.", 20, DTEAL,
     bold=True, first=True)
tb, tf = textbox(s, Inches(0.6), Inches(2.5), Inches(12.1), Inches(4.5))
bullets(tf, [
    ("Why it happens — even with modern surgery, the target is missed because of:", 0, True),
    ("Errors in biometric measurements", 1),
    ("Unpredictability of the effective lens position (ELP)", 1),
    ("Individual variation in wound healing", 1),
    ("Today's fixes are invasive — IOL exchange or corneal laser (LASIK / PRK) carry surgical risk and a long recovery.", 0),
    ("The unmet need — a way to correct sphere and cylinder non-invasively, after the eye is stable.", 0),
], size=19, gap=10)
notes(s, "Conventional fixed-power IOLs lock in a single prediction at surgery. "
         "If healing or biometry is off, the patient is left with blur, and "
         "fixing it means another operation. Residual error directly drives "
         "dissatisfaction in premium cataract surgery.")

# ========================================================================
# 3 — What is the LAL?
# ========================================================================
s = slide()
header(s, "What Is the Light Adjustable Lens?")
tb, tf = textbox(s, Inches(0.6), Inches(1.4), Inches(12.1), Inches(0.7))
para(tf, "A photoreactive silicone IOL whose dioptric power can be changed by "
         "light after implantation.", 20, DTEAL, bold=True, first=True)
tb, tf = textbox(s, Inches(0.6), Inches(2.5), Inches(6.3), Inches(4.2))
bullets(tf, [
    "Adjusted once the eye has healed and reached refractive stability",
    "Corrects both sphere and cylinder (astigmatism)",
    "Lets us customize and optimize the result for each individual patient",
], size=19, gap=14)
# history card
rect(s, Inches(7.2), Inches(2.4), Inches(5.5), Inches(3.9), fill=TEALBG)
rect(s, Inches(7.2), Inches(2.4), Pt(5), Inches(3.9), fill=TEAL)
tb, tf = textbox(s, Inches(7.55), Inches(2.6), Inches(5.0), Inches(3.5))
para(tf, "A short history", 18, DTEAL, bold=True, first=True, space_after=10)
bullets(tf, [
    "1997 — invented by Dr. Daniel Schwartz (UCSF) & Robert Grubbs (Caltech chemist, Nobel laureate)",
    "They built a lens whose 3-D structure can be altered non-invasively by light energy",
    "FDA approved for human use — Nov 22, 2017",
], size=15, gap=10, first=False)
notes(s, "Core idea: a normal-feeling monofocal-style silicone lens UNTIL we "
         "shine light on it, at which point its shape and the patient's "
         "prescription change. Grubbs won the 2005 Nobel Prize in Chemistry.")

# ========================================================================
# 4 — How It Works
# ========================================================================
s = slide()
header(s, "How It Works — Photochemistry + Diffusion")
tb, tf = textbox(s, Inches(0.6), Inches(1.4), Inches(12.1), Inches(0.6))
para(tf, "The optic holds light-sensitive macromers evenly distributed in a "
         "silicone matrix.", 19, DARK, first=True)
# flow boxes
steps = ["UV light · 365 nm\napplied to one zone", "Macromers there\npolymerize",
         "Concentration\ngradient forms", "Unexposed macromers\ndiffuse in over ~12 h",
         "Lens curvature &\npower change\n— predictably"]
bx = Inches(0.55); by = Inches(2.25); bw = Inches(2.18); bh = Inches(1.35)
gap = Inches(0.34)
from pptx.enum.shapes import MSO_SHAPE
for i, st in enumerate(steps):
    x = Emu(int(bx) + i*(int(bw)+int(gap)))
    fill = TEAL if i in (0, 4) else LIGHT
    tc = WHITE if i in (0, 4) else DARK
    sp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, by, bw, bh)
    sp.shadow.inherit = False
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    sp.line.color.rgb = DTEAL; sp.line.width = Pt(1)
    tf = sp.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    setrun(p.add_run(), st, 12.5, tc, bold=(i in (0, 4)))
    if i < 4:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                Emu(int(x)+int(bw)+Emu(int(gap)*0.1)),
                                Emu(int(by)+int(bh)//2-Inches(0.12)),
                                Emu(int(gap)*0.8), Inches(0.24))
        ar.shadow.inherit = False
        ar.fill.solid(); ar.fill.fore_color.rgb = TEAL
        ar.line.fill.background()
# three property cards
cy = Inches(4.2); cw = Inches(3.95); ch = Inches(1.7)
card(s, Inches(0.55), cy, cw, ch, LIGHT, "Targeted",
     ["Light can correct sphere and cylinder at the same time."])
card(s, Inches(4.7), cy, cw, ch, LIGHT, "Repeatable",
     ["While untreated macromers remain, the optic can be refined again."])
card(s, Inches(8.85), cy, cw, ch, LIGHT, "Predictable",
     ["Lens healing is far more uniform than corneal healing."])
notes(s, "Walk the chain left to right. Key insight: polymerising one region "
         "pulls free macromers toward it by diffusion, physically reshaping the "
         "optic. Because the material behaves consistently, the change is highly "
         "predictable — unlike the cornea.")

# ========================================================================
# 5 — Lock-in & UV Protection
# ========================================================================
s = slide()
header(s, "Lock-In & UV Protection")
tb, tf = textbox(s, Inches(0.6), Inches(1.5), Inches(6.0), Inches(3.8))
para(tf, "Locking the power in", 20, DTEAL, bold=True, first=True, space_after=12)
bullets(tf, [
    "Adjustments are repeated until surgeon and patient are satisfied",
    "The entire optic is then irradiated → all remaining macromers polymerize",
    "This \"lock-in\" permanently fixes the lens — no further change is possible",
], size=18, gap=12, first=False)
tb, tf = textbox(s, Inches(6.9), Inches(1.5), Inches(5.8), Inches(3.8))
para(tf, "Guarding against stray UV", 20, DTEAL, bold=True, first=True, space_after=12)
bullets(tf, [
    "ActivShield — an integrated UV-absorbing layer (2nd generation, 2021)",
    "Prevents accidental sunlight from altering the lens before lock-in",
    "Patients wear UV-protective glasses from surgery until 24 h after the final treatment",
], size=18, gap=12, first=False)
rect(s, Inches(0.6), Inches(5.7), Inches(12.1), Inches(0.9), fill=TEALBG)
tb, tf = textbox(s, Inches(0.6), Inches(5.7), Inches(12.1), Inches(0.9),
                 anchor=MSO_ANCHOR.MIDDLE)
para(tf, "Until lock-in, the lens is a \"living\" optic — protection is essential.",
     18, DTEAL, bold=True, first=True, align=PP_ALIGN.CENTER)
notes(s, "Lock-in is the irreversible commitment step. Until you reach it, "
         "ambient UV could polymerise macromers uncontrollably. ActivShield plus "
         "the goggles make the adjustment window safe. Glasses compliance is "
         "non-negotiable.")

# ========================================================================
# 6 — Lens Specifications
# ========================================================================
s = slide()
header(s, "Lens Specifications")
tb, tf = textbox(s, Inches(0.6), Inches(1.4), Inches(6.0), Inches(0.5))
para(tf, "Optic", 19, DTEAL, bold=True, first=True)
optic = [("Material", "Photoreactive, UV-absorbing silicone"),
         ("Refractive index", "1.43"),
         ("Diopter range", "+10.0 to +30.0 D"),
         ("Edge design", "Rounded anterior · square posterior"),
         ("Optic diameter", "6 mm"),
         ("Overall diameter", "13 mm")]
def spec_table(s, x, y, w, rows, hl_idx=None):
    t = s.shapes.add_table(len(rows), 2, x, y, w, Inches(0.5*len(rows))).table
    t.columns[0].width = int(w*0.42); t.columns[1].width = int(w*0.58)
    for r, (k, v) in enumerate(rows):
        for c, val in enumerate((k, v)):
            cell = t.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHT if r % 2 == 0 else WHITE
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_top = Pt(2); cell.margin_bottom = Pt(2)
            cell.margin_left = Pt(8)
            p = cell.text_frame.paragraphs[0]
            bold = (c == 0) or (hl_idx is not None and r == hl_idx)
            setrun(p.add_run(), val, 13.5,
                   DTEAL if (hl_idx is not None and r == hl_idx and c == 1) else DARK,
                   bold=bold)
    return t
spec_table(s, Inches(0.6), Inches(1.95), Inches(6.0), optic, hl_idx=4)
tb, tf = textbox(s, Inches(7.1), Inches(1.4), Inches(5.6), Inches(0.5))
para(tf, "Haptics", 19, DTEAL, bold=True, first=True)
spec_table(s, Inches(7.1), Inches(1.95), Inches(5.6),
           [("Material", "Blue-core PMMA monofilament"), ("Haptic angle", "10°")])
rect(s, Inches(7.1), Inches(3.4), Inches(5.6), Inches(1.6), fill=AMBBG)
rect(s, Inches(7.1), Inches(3.4), Pt(5), Inches(1.6), fill=AMBER)
tb, tf = textbox(s, Inches(7.4), Inches(3.55), Inches(5.1), Inches(1.3),
                 anchor=MSO_ANCHOR.MIDDLE)
para(tf, "A densely UV-filtering layer in the posterior optic shields the "
         "retina during every light treatment and lock-in.", 14, AMBER,
     bold=True, first=True)
notes(s, "Don't read the whole table — highlight that it's a silicone lens with "
         "a 6 mm optic (matters for dilation later) and a built-in posterior UV "
         "filter to protect the macula. Diopter steps are finer (0.5 D) in the "
         "common +16 to +24 range.")

# ========================================================================
# 7 — Surgical Workflow
# ========================================================================
s = slide()
header(s, "The Surgical Workflow")
wf = ["1 · Pre-op biometry\nbase monofocal power", "2 · Standard\ncataract surgery",
      "3 · ~2–4 weeks\nrefraction +\ncustom prescription",
      "4 · LDD light\nadjustments\n2–4 sessions · 365 nm", "5 · Lock-in\npower fixed"]
bx = Inches(0.55); by = Inches(1.9); bw = Inches(2.18); bh = Inches(1.5); gap = Inches(0.34)
for i, st in enumerate(wf):
    x = Emu(int(bx) + i*(int(bw)+int(gap)))
    fill = NAVY if i == 0 else (TEAL if i == 4 else LIGHT)
    tc = WHITE if i in (0, 4) else DARK
    sp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, by, bw, bh)
    sp.shadow.inherit = False
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    sp.line.color.rgb = DTEAL; sp.line.width = Pt(1)
    tf = sp.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    setrun(p.add_run(), st, 12.5, tc, bold=(i in (0, 4)))
    if i < 4:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                Emu(int(x)+int(bw)+Emu(int(gap)*0.1)),
                                Emu(int(by)+int(bh)//2-Inches(0.12)),
                                Emu(int(gap)*0.8), Inches(0.24))
        ar.shadow.inherit = False
        ar.fill.solid(); ar.fill.fore_color.rgb = TEAL; ar.line.fill.background()
card(s, Inches(0.55), Inches(4.0), Inches(5.9), Inches(1.9), LIGHT,
     "The Light Delivery Device (LDD)",
     ["An optical projection system + UV source on a slit lamp;",
      "light is focused onto the lens through a special corneal",
      "contact lens."])
rect(s, Inches(6.8), Inches(4.0), Inches(5.9), Inches(1.9), fill=AMBBG)
rect(s, Inches(6.8), Inches(4.0), Pt(5), Inches(1.9), fill=AMBER)
tb, tf = textbox(s, Inches(7.1), Inches(4.15), Inches(5.4), Inches(1.6))
para(tf, "⚠ Needs ~7 mm dilation", 15, AMBER, bold=True, first=True, space_after=6)
para(tf, "The optic is 6 mm, so the pupil must open wide enough for UV light to "
         "cover the entire optical zone.", 13, DARK)
notes(s, "Stress the timeline: patient leaves surgery 'undialed', we wait for "
         "stability (~2–4 wks), sculpt the prescription over 2–4 visits, then "
         "lock in. Whole course usually done by ~5–6 weeks. Dilation requirement "
         "is a real gatekeeper.")

# ========================================================================
# 8 — Indications & Contraindications
# ========================================================================
s = slide()
header(s, "Indications & Contraindications")
rect(s, Inches(0.6), Inches(1.4), Inches(6.0), Inches(0.55), fill=TEALBG)
tb, tf = textbox(s, Inches(0.6), Inches(1.4), Inches(6.0), Inches(0.55),
                 anchor=MSO_ANCHOR.MIDDLE)
para(tf, "✓  Good candidates", 18, DTEAL, bold=True, first=True)
tb, tf = textbox(s, Inches(0.6), Inches(2.1), Inches(6.0), Inches(4.8))
bullets(tf, [
    "High risk of refractive surprise — prior LASIK, PRK or RK, where IOL calculation is uncertain",
    "Premium expectations without diffractive multifocals — works on a monofocal principle, target set post-op",
    "Mini-monovision / blended vision — tested and tuned after surgery",
    "Low residual astigmatism — correctable down to ~0.50 D",
], size=16, gap=12, first=True)
rect(s, Inches(7.0), Inches(1.4), Inches(5.7), Inches(0.55), fill=AMBBG)
tb, tf = textbox(s, Inches(7.0), Inches(1.4), Inches(5.7), Inches(0.55),
                 anchor=MSO_ANCHOR.MIDDLE)
para(tf, "⛔  Avoid in", 18, AMBER, bold=True, first=True)
tb, tf = textbox(s, Inches(7.0), Inches(2.1), Inches(5.7), Inches(4.8))
bullets(tf, [
    "Pre-existing macular disease",
    "Prior ocular herpes infection",
    "UV-sensitizing drugs, or retinotoxic drugs (e.g. Tamoxifen)",
    "Nystagmus",
    "Patients who cannot comply with the adjustment schedule or UV-glasses regimen",
], size=16, gap=11, first=True)
notes(s, "Frame indications around the headline use case: post-refractive eyes, "
         "where everyone struggles with IOL power. Contraindications cluster into "
         "(1) UV/retinal risk and (2) ability to deliver/comply with treatment.")

# ========================================================================
# 9 — Advantage: Post-Refractive Surgery Eyes
# ========================================================================
s = slide()
header(s, "Advantage — Post-Refractive Surgery Eyes")
tb, tf = textbox(s, Inches(0.6), Inches(1.4), Inches(12.1), Inches(0.7))
para(tf, "After LASIK / PRK / RK, IOL power calculation is unreliable — the LAL "
         "corrects the surprise non-invasively.", 19, DTEAL, bold=True, first=True)
# stat card 1
rect(s, Inches(0.6), Inches(2.35), Inches(5.9), Inches(2.0), fill=LIGHT)
tb, tf = textbox(s, Inches(0.85), Inches(2.5), Inches(5.4), Inches(1.8))
para(tf, "Brierley — 34 eyes, prior refractive surgery", 15, DARK, bold=True,
     first=True, space_after=4)
para(tf, "97%  ·  100%", 34, TEAL, bold=True, space_after=2)
para(tf, "within ±0.50 D  ·  within ±1.00 D of target", 14, DARK, space_after=2)
para(tf, "≈ 60% more predictable than monofocal IOLs in these eyes", 12, GRAY,
     italic=True)
# stat card 2
rect(s, Inches(6.8), Inches(2.35), Inches(5.9), Inches(2.0), fill=LIGHT)
tb, tf = textbox(s, Inches(7.05), Inches(2.5), Inches(5.4), Inches(1.8))
para(tf, "Wong & Folden — 2nd-gen LAL, post-LASIK/PRK", 15, DARK, bold=True,
     first=True, space_after=4)
para(tf, "82%  ·  97%", 34, TEAL, bold=True, space_after=2)
para(tf, "saw 20/20+  ·  within ±0.50 D (mean SE 0.01 D)", 14, DARK)
rect(s, Inches(0.6), Inches(4.7), Inches(12.1), Inches(1.5), fill=TEALBG)
tb, tf = textbox(s, Inches(0.9), Inches(4.85), Inches(11.5), Inches(1.2),
                 anchor=MSO_ANCHOR.MIDDLE)
para(tf, "vs. corneal laser enhancement", 16, DTEAL, bold=True, first=True,
     space_after=4)
para(tf, "Adjustment can begin at ~3 weeks (not the ~3-month wait for corneal "
         "stabilization), is more predictable, non-invasive, and improves "
         "patient satisfaction.", 15, DARK)
notes(s, "Strongest clinical story. Post-refractive eyes are the classic "
         "'refractive surprise' population because the altered cornea breaks "
         "standard IOL formulas. The LAL sidesteps this by measuring the actual "
         "healed refraction. Headline: Brierley's ~60% predictability gain.")

# ========================================================================
# 10 — Clinical Outcomes vs Other IOLs
# ========================================================================
s = slide()
header(s, "Clinical Outcomes vs. Other IOLs")
tb, tf = textbox(s, Inches(0.6), Inches(1.35), Inches(6.0), Inches(0.6))
para(tf, "Monocular UCDVA ≥ 20/20", 17, DTEAL, bold=True, first=True, space_after=0)
para(tf, "(Nakagama & Doane — 150 eyes)", 12, GRAY, italic=True)
t = s.shapes.add_table(4, 2, Inches(0.6), Inches(2.25), Inches(5.9),
                       Inches(1.9)).table
t.columns[0].width = Inches(3.7); t.columns[1].width = Inches(2.2)
rows = [("Lens", "20/20 or better"), ("LAL", "64 %"),
        ("Toric monofocal", "46 %"), ("Spherical monofocal", "32 %")]
for r, (a, b) in enumerate(rows):
    head = (r == 0); hl = (r == 1)
    for c, val in enumerate((a, b)):
        cell = t.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY if head else (TEALBG if hl else (LIGHT if r % 2 == 0 else WHITE))
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE; cell.margin_left = Pt(8)
        p = cell.text_frame.paragraphs[0]
        setrun(p.add_run(), val, 14,
               WHITE if head else (DTEAL if hl else DARK), bold=head or hl)
tb, tf = textbox(s, Inches(0.6), Inches(4.35), Inches(5.9), Inches(2.6))
para(tf, "Within ±0.50 D of plano:  LAL 92% vs toric 82% vs spherical 64%.",
     14, DARK, first=True, space_after=8)
para(tf, "FDA trial (600 eyes): 70% of LAL reached 20/20 vs 36% of monofocal "
         "controls.", 14, DARK)
# right column
rect(s, Inches(6.9), Inches(1.55), Inches(5.8), Inches(2.0), fill=LIGHT)
tb, tf = textbox(s, Inches(7.15), Inches(1.7), Inches(5.3), Inches(1.8),
                 anchor=MSO_ANCHOR.MIDDLE)
para(tf, "Customized monovision", 16, DARK, bold=True, first=True, space_after=2)
para(tf, "96%", 40, TEAL, bold=True, space_after=0)
para(tf, "achieved binocular 20/20 distance + J2 near (Folden & Wong)", 13, DARK)
rect(s, Inches(6.9), Inches(3.8), Inches(5.8), Inches(2.4), fill=TEALBG)
rect(s, Inches(6.9), Inches(3.8), Pt(5), Inches(2.4), fill=TEAL)
tb, tf = textbox(s, Inches(7.2), Inches(3.95), Inches(5.3), Inches(2.1),
                 anchor=MSO_ANCHOR.MIDDLE)
para(tf, "LAL+ (2023)", 16, DTEAL, bold=True, first=True, space_after=4)
para(tf, "Adds central power for a broader depth of focus: better intermediate "
         "& near vision and less anisometropia, with no loss of distance acuity.",
     14, DARK)
notes(s, "The LAL roughly doubles the rate of 20/20 uncorrected vision vs a "
         "spherical monofocal, and beats toric on residual error too — because "
         "it's tuned to the real post-op refraction. Monovision: patients trial "
         "the blur before committing. LAL+ widens the range of focus.")

# ========================================================================
# 11 — Limitations & Safety
# ========================================================================
s = slide()
header(s, "Limitations & Safety")
rect(s, Inches(0.6), Inches(1.4), Inches(6.0), Inches(0.55), fill=AMBBG)
tb, tf = textbox(s, Inches(0.6), Inches(1.4), Inches(6.0), Inches(0.55),
                 anchor=MSO_ANCHOR.MIDDLE)
para(tf, "Limitations", 18, AMBER, bold=True, first=True)
tb, tf = textbox(s, Inches(0.6), Inches(2.1), Inches(6.0), Inches(4.0))
bullets(tf, [
    "Premium cost — plus added service costs for refractions, light sessions, lock-ins",
    "Needs adequate dilation (6.5–7 mm)",
    "Corneal astigmatism may drift with age, affecting long-term result",
    "Heavily compliance-dependent (UV glasses, multiple visits)",
], size=16, gap=12, first=True)
rect(s, Inches(7.0), Inches(1.4), Inches(5.7), Inches(0.55), fill=TEALBG)
tb, tf = textbox(s, Inches(7.0), Inches(1.4), Inches(5.7), Inches(0.55),
                 anchor=MSO_ANCHOR.MIDDLE)
para(tf, "Safety profile", 18, DTEAL, bold=True, first=True)
tb, tf = textbox(s, Inches(7.0), Inches(2.1), Inches(5.7), Inches(4.0))
bullets(tf, [
    "Overall similar to standard cataract surgery",
    "LAL-specific effects are mostly transient and UV-related (e.g. temporary red-tinged vision)",
    "Secondary surgical intervention: only 1.7% (FDA study)",
], size=16, gap=12, first=True)
tb, tf = textbox(s, Inches(0.6), Inches(6.4), Inches(12.1), Inches(0.7))
para(tf, "Future direction: pairing the LAL with adaptive-optics to correct "
         "aberrations in real time.", 14, GRAY, italic=True, first=True)
notes(s, "Be balanced. Big practical downsides are cost and the logistics of "
         "multiple visits, both compliance-dependent. Safety-wise it behaves like "
         "routine cataract surgery; UV side effects are almost always temporary; "
         "further surgery is rare at 1.7%.")

# ========================================================================
# 12 — Key Takeaways
# ========================================================================
s = slide()
rect(s, 0, 0, SW, SH, fill=WHITE)
rect(s, 0, 0, SW, Inches(1.15), fill=NAVY)
rect(s, 0, Inches(1.15), SW, Pt(4), fill=TEAL)
tb, tf = textbox(s, Inches(0.6), Inches(0.12), Inches(12.1), Inches(0.95),
                 anchor=MSO_ANCHOR.MIDDLE)
para(tf, "Key Takeaways", 30, WHITE, bold=True, first=True)
takeaways = [
    ("The first truly adjustable IOL — non-invasive refractive fine-tuning after surgery", TEALBG, DTEAL),
    ("Excellent predictability, especially in post-refractive-surgery eyes", TEALBG, DTEAL),
    ("Enables customized monovision / blended vision patients can trial first", TEALBG, DTEAL),
    ("Trade-offs — cost, multiple visits, UV-glasses compliance, adequate dilation", AMBBG, AMBER),
    ("Safe, with promising and stable refractive outcomes", TEALBG, DTEAL),
]
y = Inches(1.55)
for txt, bg, accent in takeaways:
    rect(s, Inches(1.4), y, Inches(10.5), Inches(0.92), fill=bg)
    rect(s, Inches(1.4), y, Pt(6), Inches(0.92), fill=accent)
    tb, tf = textbox(s, Inches(1.75), y, Inches(10.0), Inches(0.92),
                     anchor=MSO_ANCHOR.MIDDLE)
    para(tf, txt, 17, DARK, bold=True, first=True)
    y = Emu(int(y) + int(Inches(1.07)))
notes(s, "Land the plane: the LAL converts cataract surgery from a one-shot "
         "prediction into an adjustable, patient-in-the-loop refractive process. "
         "Shines where prediction is hardest (post-refractive) and where "
         "personalization matters (monovision). Price is cost and compliance.")

# ========================================================================
# 13 — Thank You
# ========================================================================
s = slide()
rect(s, 0, 0, SW, SH, fill=NAVY)
rect(s, 0, Inches(4.0), SW, Pt(4), fill=TEAL)
tb, tf = textbox(s, Inches(1.0), Inches(2.6), Inches(11.3), Inches(1.6),
                 anchor=MSO_ANCHOR.MIDDLE)
para(tf, "Thank You", 48, WHITE, bold=True, first=True, align=PP_ALIGN.CENTER,
     space_after=6)
para(tf, "Questions?", 24, RGBColor(0x9D, 0xD9, 0xD2), align=PP_ALIGN.CENTER)
tb, tf = textbox(s, Inches(1.0), Inches(5.7), Inches(11.3), Inches(1.3),
                 anchor=MSO_ANCHOR.MIDDLE)
para(tf, "Sources: Word manuscript (primary) · Jun et al., Curr Opin Ophthalmol "
         "2024 · Doane et al., JCRS 2025 · Wong & Folden, Clin Ophthalmol 2023 · "
         "Nakagama & Doane, Missouri Medicine 2025", 12,
     RGBColor(0xCB, 0xE7, 0xE3), italic=True, first=True, align=PP_ALIGN.CENTER)

prs.save("lal-presentation-editable.pptx")
print("saved lal-presentation-editable.pptx |", len(prs.slides._sldIdLst), "slides")
